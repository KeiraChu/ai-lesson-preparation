import io
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation


def extract_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension in {".txt", ".md"}:
        return content.decode("utf-8", errors="replace")
    if extension == ".pdf":
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages)
    if extension == ".docx":
        document = Document(io.BytesIO(content))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    if extension == ".pptx":
        presentation = Presentation(io.BytesIO(content))
        return "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )
    if extension == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return "\n".join(
            "\t".join("" if value is None else str(value) for value in row)
            for sheet in workbook.worksheets for row in sheet.iter_rows(values_only=True)
        )
    raise ValueError(f"unsupported file type: {extension}")
